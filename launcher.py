"""One-click entry point for the portable Windows app and source checkout."""
import asyncio
import getpass
import os
from pathlib import Path
import re
import sys

from paths import DATA_DIR


def save_token(token, path):
    if not re.fullmatch(r'\d+:[A-Za-z0-9_-]{20,}', token):
        raise ValueError('That does not look like a Telegram bot token. Copy the complete token from @BotFather.')
    path.parent.mkdir(parents=True,exist_ok=True)
    # Atomic update; token is never printed or passed as a process argument.
    temporary = path.with_suffix('.tmp')
    existing=path.read_text(encoding='utf-8').splitlines() if path.exists() else []
    lines=[line for line in existing if not re.match(r'^\s*(?:export\s+)?BOT_TOKEN\s*=',line)]
    temporary.write_text('\n'.join(lines+['BOT_TOKEN='+token])+'\n', encoding='utf-8')
    temporary.replace(path)


async def check_token(token):
    from runtime import build_app
    app = build_app(token,'setup')
    async with app.bot:
        return app.bot.username


def configure(force=False):
    from dotenv import dotenv_values
    path = DATA_DIR / '.env'
    if not force and (os.getenv('BOT_TOKEN') or path.exists() and dotenv_values(path).get('BOT_TOKEN')):
        return
    print('\nFIRST-TIME SETUP (once per computer)\n')
    print('1. Open https://t.me/BotFather in Telegram.')
    print('2. Send /newbot and copy the token. For an existing bot, reuse its token.')
    print('3. Paste it below. The token stays on this computer.\n')
    token = getpass.getpass('Bot token (hidden): ').strip()
    if not re.fullmatch(r'\d+:[A-Za-z0-9_-]{20,}', token):
        raise ValueError('Invalid token format. Restart and paste the complete token.')
    print('Checking your Telegram connection...')
    try:
        username = asyncio.run(check_token(token))
    except Exception:
        raise RuntimeError('Could not verify the token. Check your connection and the token, then try again. No token was saved.') from None
    save_token(token,path)
    os.environ['BOT_TOKEN'] = token
    print(f'Connected to @{username}. Setup saved.\n')


def self_check():
    """Offline smoke test of the compiled executable; never touches real user data."""
    import tempfile
    import database_sqlite as db
    from charts import chart
    from diary import quick_parse
    from openpyxl import Workbook, load_workbook
    from datetime import date
    import io
    from runtime import build_app
    original = db.DB_PATH
    with tempfile.TemporaryDirectory() as tmp:
        try:
            db.DB_PATH = Path(tmp)/'smoke.sqlite3'
            db.ensure_bot_schema()
            t = db.create_training(1)
            db.add_exercise_to_training(t['training_id'],quick_parse('Жим лежачи 3x10 60 кг'))
            assert len(db.get_training_exercises(t['training_id'])[0]['sets']) == 3
            png = chart([dict(kind='training',name='Test',day=date.today().isoformat(),sets=3,reps=10,value=60,unit='кг')],7)
            assert png.startswith(b'\x89PNG')
            stream = io.BytesIO()
            Workbook().save(stream)
            assert load_workbook(io.BytesIO(stream.getvalue())).active is not None
            build_app('123456:AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA','test')
            from coach_rag import library, retrieve, references
            assert len(library()['cards']) >= 7 and len(references())==5
            assert retrieve('гіпертрофія обсяг')
            from unittest.mock import patch
            import backups
            with patch.object(backups,'DATA_DIR',Path(tmp)):
                snapshot=backups.create_backup()
                db.set_setting('smoke_marker','temporary')
                backups.restore_backup(snapshot)
                assert db.setting('smoke_marker') is None
        finally:
            db.DB_PATH = original
    print('SELF-CHECK PASSED: database, parser, charts, Excel, backup/restore, local evidence retrieval and Telegram application. No LLM started.')


def main():
    for stream in (sys.stdout,sys.stderr):
        if hasattr(stream,'reconfigure'):
            stream.reconfigure(encoding='utf-8',errors='replace')
    os.environ.setdefault('MPLCONFIGDIR',str(DATA_DIR/'cache'/'matplotlib'))
    if '--check' in sys.argv:
        self_check()
        return
    if '--stop' in sys.argv:
        DATA_DIR.mkdir(parents=True,exist_ok=True)
        (DATA_DIR/'stop.request').touch()
        print('Stop requested. The bot will shut down within a few seconds.')
        return
    print('\n  NEXTSET UA\n  Your private training diary in Telegram\n' + '  ' + '-'*43)
    print(f'  Data folder: {DATA_DIR}\n')
    try:
        if '--restore' in sys.argv:
            from backups import restore_console
            restore_console()
            return
        if '--evaluate-ai' in sys.argv:
            from dotenv import load_dotenv
            load_dotenv(DATA_DIR/'.env')
            from ai_evaluation import main as evaluate
            evaluate()
            return
        if '--setup-ai' in sys.argv:
            from ai_setup import setup
            setup()
            if sys.stdin and sys.stdin.isatty():
                input('\nPress Enter to close...')
            return
        configure(force='--setup' in sys.argv)
        from dotenv import load_dotenv
        load_dotenv(DATA_DIR/'.env')
        from ai_setup import start_if_enabled
        start_if_enabled()
        from runtime import main as run_bot
        run_bot()
    except KeyboardInterrupt:
        print('\nStopped. Your saved workouts are safe.')
    except (RuntimeError,ValueError) as error:
        print('\n' + str(error))
    except Exception:
        print('\nCould not start. Check internet access, folder permissions and your bot token.')
        print('If this bot is running on another PC, stop it there first.')
    if sys.stdin and sys.stdin.isatty():
        input('\nPress Enter to close...')


if __name__ == '__main__':
    main()
