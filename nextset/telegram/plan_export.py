"""On-demand Excel plan export using the application's existing Excel dependency."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from nextset.coaching.presentation import schedule, span
from nextset.core.exercises import NAMES


def workbook_bytes(block, state, today=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'План тренувань'
    sheet.append([f'План на {block["length"]} тижнів — ' + ('прийнято' if block.get('approved') else 'проєкт')])
    sheet.append(['Кг лише для поточної розрахованої сесії; підтвердь її перед тренуванням. Порожньо = ще не розраховано. Дати орієнтовні.'])
    sheet.merge_cells('A1:F1')
    sheet.merge_cells('A2:F2')
    sheet.append(['Дата', 'Вправа', 'Підходи', 'Повторення', 'Кг', 'Запас повторень'])
    bands = [1, 3]
    for slot, items in schedule(block, state, today):
        if slot['session'] == 0:
            sheet.append([f'Тиждень {slot["week"]}'])
            sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6)
            bands.append(sheet.max_row)
            sheet.append([block['weeks'][slot['week'] - 1]])
            sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6)
            sheet.row_dimensions[sheet.max_row].height = 45
        sheet.append([f'{chr(65 + slot["session"])}. {slot["title"]}'])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6)
        bands.append(sheet.max_row)
        for item in items:
            sheet.append([slot['date'], NAMES[item['exercise']], span(item['sets']), span(item['reps']), item.get('weight'), span(item['rir'])])
            sheet.cell(sheet.max_row, 1).number_format = 'dd.mm.yyyy'
            sheet.cell(sheet.max_row, 5).number_format = '0.##'
            sheet.row_dimensions[sheet.max_row].height = 32
    for row in sheet:
        for cell in row:
            # Model-authored titles/guidance are literal text, never formulas.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11)
            if cell.row in bands:
                cell.fill = PatternFill('solid', fgColor='DCE6F1')
                cell.font = Font(name='Calibri', size=11, bold=True)
    for column, width in zip('ABCDEF', [15, 43, 14, 16, 12, 20]):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 34
    sheet.freeze_panes = 'C4'
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:3'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    book.save(output)
    return output.getvalue()
