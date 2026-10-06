import asyncio
from contextlib import redirect_stdout
import io
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
import database_sqlite as db
import diary
import local_ai
import handlers_training as training
import handlers_exercises as exercises
from handlers_export import generate_csv_export, generate_excel_report
from utils_constants import *
from telegram.ext import ApplicationHandlerStop


class Tests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.original=db.DB_PATH
        db.DB_PATH=Path(self.temp.name)/'test.sqlite3'
        db.ensure_bot_schema()
        self.uid=42
        self.msg=SimpleNamespace(text='',message_id=10,from_user=SimpleNamespace(id=42,username='test',first_name='Тест'),reply_text=AsyncMock(),reply_photo=AsyncMock(),reply_document=AsyncMock())
        self.update=SimpleNamespace(message=self.msg,effective_message=self.msg,effective_user=self.msg.from_user,effective_chat=SimpleNamespace(type='private'))
        self.ctx=SimpleNamespace(user_data={},bot=SimpleNamespace(send_message=AsyncMock()),application=SimpleNamespace(bot_data={'pair_code':'pair'}))

    def tearDown(self):
        db.DB_PATH=self.original
        self.temp.cleanup()

    def seed(self):
        t=db.create_training(self.uid)
        e=diary.quick_parse('Жим лежачи 3x10 60 кг')
        db.add_exercise_to_training(t['training_id'],e)
        return t

    async def test_guided_session_and_history(self):
        self.assertEqual(await training.start_training(self.update,self.ctx),INPUT_MEASUREMENTS_CHOICE)
        self.msg.text='⏭️ Пропустити виміри'
        self.assertEqual(await training.handle_measurements_choice(self.update,self.ctx),TRAINING_MENU)
        self.msg.text='Жим лежачи'
        self.assertEqual(await training.handle_strength_exercise_selection(self.update,self.ctx),INPUT_SETS)
        self.msg.text='60 10\n62,5 8\n65 6'
        await training.handle_set_input(self.update,self.ctx)
        await training.save_exercise(self.update,self.ctx)
        rows=db.get_current_training(self.uid)['exercises']
        self.assertEqual([s['weight'] for s in rows[0]['sets']],[60,62.5,65])
        await training.show_finish_summary(self.update,self.ctx)
        self.msg.text='✅ Підтвердити завершення'
        await training.handle_finish_confirmation(self.update,self.ctx)
        self.assertIsNone(db.get_current_training(self.uid))
        await diary.send_history(self.update,self.ctx)
        text=self.msg.reply_text.call_args.args[0]
        self.assertIn('65 кг × 6',text)
        self.assertIn('Жим лежачи',text)

    async def test_invalid_batch_atomic(self):
        await training.start_training(self.update,self.ctx)
        self.msg.text='Жим лежачи'
        await training.handle_strength_exercise_selection(self.update,self.ctx)
        for text in ['60 10\n-3 8','nan 10','60 0','60 10 extra']:
            self.msg.text=text
            await training.handle_set_input(self.update,self.ctx)
            self.assertEqual(self.ctx.user_data['current_exercise']['sets'],[])

    async def test_quick_dedup(self):
        e=diary.quick_parse('Жим лежачи 3х10 62,5 кг')
        await diary.save_quick(self.update,self.ctx,e)
        await diary.save_quick(self.update,self.ctx,e)
        t=db.get_current_training(self.uid)
        self.assertEqual(len(t['exercises']),1)
        self.assertEqual(t['exercises'][0]['sets'][0]['weight'],62.5)

    async def test_measurements_validation_and_history(self):
        await training.start_training(self.update,self.ctx)
        self.msg.text='біципс 0 см'
        self.assertEqual(await training.save_measurements(self.update,self.ctx),INPUT_MEASUREMENTS)
        self.msg.text='біцепс лівий 35,5 см; вага 80 кг'
        self.assertEqual(await training.save_measurements(self.update,self.ctx),TRAINING_MENU)
        self.assertEqual(len(db.get_measurements_history(self.uid)),1)
        rows=diary.graph_rows(self.uid,30,'біцепс лівий')
        self.assertEqual(rows[0]['value'],35.5)

    async def test_measurement_back_button(self):
        await training.start_training(self.update,self.ctx)
        self.msg.text='🔙 Головне меню'
        self.assertEqual(await training.handle_measurements_choice(self.update,self.ctx),MAIN_MENU)

    async def test_custom_exercise_add_delete(self):
        self.msg.text='💪 Силова вправа'
        await exercises.add_custom_exercise_mgmt(self.update,self.ctx)
        self.msg.text='Тестова вправа'
        await exercises.save_new_strength_exercise_mgmt(self.update,self.ctx)
        self.assertIn('Тестова вправа',db.get_visible_exercise_lists(self.uid)['strength'])
        self.msg.text='💪 Тестова вправа'
        await exercises.delete_exercise_handler(self.update,self.ctx)
        self.assertNotIn('Тестова вправа',db.get_visible_exercise_lists(self.uid)['strength'])

    async def test_private_pairing(self):
        from runtime import guard
        self.msg.text='/start wrong'
        with self.assertRaises(ApplicationHandlerStop):
            await guard(self.update,self.ctx)
        self.msg.text='/start pair'
        self.update.effective_chat.type='group'
        with self.assertRaises(ApplicationHandlerStop):
            await guard(self.update,self.ctx)
        self.update.effective_chat.type='private'
        # Redirected Windows output may use cp1252 rather than UTF-8.
        with io.TextIOWrapper(io.BytesIO(),encoding='cp1252',errors='strict') as console:
            with redirect_stdout(console):
                await guard(self.update,self.ctx)
        self.assertEqual(db.setting('owner'),42)
        self.update.effective_user=SimpleNamespace(id=43)
        with self.assertRaises(ApplicationHandlerStop):
            await guard(self.update,self.ctx)

    async def test_graphs_and_ranges(self):
        t=self.seed()
        rows=diary.graph_rows(self.uid,7)
        self.assertEqual(len(rows),3)
        with db.connection() as c:
            c.execute('UPDATE trainings SET date_start=?',( (datetime.now()-timedelta(days=7)).isoformat(),))
        self.assertEqual(diary.graph_rows(self.uid,7),[])
        self.assertEqual(len(diary.graph_rows(self.uid,30)),3)
        await diary.show_graph(self.update,self.ctx,'30 Жим лежачи')
        self.assertTrue(self.msg.reply_photo.called)

    async def test_export(self):
        t=self.seed()
        db.finish_training(t['training_id'])
        csv=generate_csv_export(self.uid)
        self.assertIn('Жим лежачи',csv)
        self.assertIn('Вага (кг)',csv)
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(generate_excel_report(self.uid,'all_time')))
        self.assertIn('Подробиці підходів',wb.sheetnames)
        self.assertEqual(wb['Подробиці підходів'].max_row,4)

    async def test_owner_data_isolation(self):
        t=self.seed()
        eid=db.get_training_exercises(t['training_id'])[0]['exercise_id']
        self.assertFalse(db.delete_exercise_record(99,eid))
        self.assertEqual(db.get_user_trainings(99,include_active=True),[])
        self.assertTrue(db.delete_exercise_record(42,eid))

    async def test_ai_proposal_never_saves_without_confirmation(self):
        proposed=diary.quick_parse('Жим лежачи 3x10 60 кг')
        self.msg.text='/ai сьогодні жим 3 по 10 на 60 кг'
        with patch('local_ai.propose',return_value=proposed):
            with self.assertRaises(ApplicationHandlerStop):
                await diary.extras(self.update,self.ctx)
        self.assertIsNone(db.get_current_training(42))
        self.msg.text='/confirm'
        with self.assertRaises(ApplicationHandlerStop):
            await diary.extras(self.update,self.ctx)
        self.assertEqual(len(db.get_current_training(42)['exercises']),1)

    async def test_bad_ai_output(self):
        for value in [{},{'unclear':True},{'unclear':False,'name':'Жим','sets':[{'weight':float('nan'),'reps':8}]},{'unclear':False,'name':'Жим','sets':[{'weight':60,'reps':True}]}]:
            with self.assertRaises(ValueError):
                local_ai.validate(value)

    async def test_cardio_invalid(self):
        await training.start_training(self.update,self.ctx)
        self.ctx.user_data['current_exercise']={'name':'Біг','type':'cardio'}
        self.msg.text='-5 10'
        self.assertEqual(await training.save_cardio_exercise(self.update,self.ctx,'km_h'),INPUT_CARDIO_KM_H)
        self.assertEqual(db.get_current_training(42)['exercises'],[])

    async def test_reminders_once_per_day(self):
        self.seed()
        db.set_setting('owner',42)
        with db.connection() as c:
            c.execute('UPDATE trainings SET date_start=?',('2026-09-01T12:00:00',))
        fixed=datetime(2026,10,5,19,0)
        with patch('diary.datetime') as clock:
            clock.now.return_value=fixed
            await diary.reminder_job(self.ctx)
            self.ctx.bot.send_message.assert_not_called()
            db.set_setting('reminder_days',10)
            await diary.reminder_job(self.ctx)
            await diary.reminder_job(self.ctx)
            self.ctx.bot.send_message.assert_awaited_once()

    async def test_export_formula_sanitization(self):
        t=db.create_training(42)
        db.add_exercise_to_training(t['training_id'],{'name':'=1+1','type':'strength','sets':[{'weight':10,'reps':8}]})
        db.finish_training(t['training_id'])
        self.assertIn("'=1+1",generate_csv_export(42))

if __name__=='__main__':
    unittest.main()
