import copy
from datetime import date
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook
from nextset.coaching.presentation import overview, session_items
from nextset.telegram.plan_export import workbook_bytes
from nextset.telegram import ui


def fixture():
    block = dict(length=8, approved=True, created='block-1', start_cursor=0,
                 schedule_start='2026-10-19', weeks=['Рекомендація'] * 8,
                 sessions=[dict(title=f'Тренування {i}', items=[dict(exercise='bench', sets=[3, 4], reps=[8, 12], rir=[2, 3])]) for i in range(3)])
    state = dict(cursor=0, block=block, pending=dict(block_id='block-1', cursor=0,
                 day='2026-10-19', bonus=False, items=[dict(key='bench', sets=3, reps=8, rir=2, weight=60)]))
    return block, state


class PlanPresentation(unittest.TestCase):
    def test_weights_only_in_matching_current_session(self):
        block, state = fixture()
        today = date(2026, 10, 19)
        self.assertEqual(session_items(block, state, 0, today)[0]['weight'], 60)
        self.assertIsNone(session_items(block, state, 3, today)[0]['weight'])
        for change in ({'day': '2026-10-18'}, {'block_id': 'old'}, {'bonus': True}):
            other = copy.deepcopy(state)
            other['pending'].update(change)
            self.assertIsNone(session_items(block, other, 0, today)[0]['weight'])
        block['approved'] = False
        self.assertIsNone(session_items(block, state, 0, today)[0]['weight'])

    def test_substitutions_and_unknown_weight(self):
        block, state = fixture()
        block['overrides'] = {'1': {'bench': 'squat'}}
        self.assertEqual(session_items(block, state, 1)[0]['exercise'], 'squat')
        text = '\n'.join(overview(block, state, date(2026, 10, 19)))
        self.assertIn('60 кг', text)
        self.assertIn('ще не розраховано', text)

    def test_excel_all_sessions_numeric_kg_no_formula_injection(self):
        block, state = fixture()
        block['weeks'][0] = '=HYPERLINK("https://example.com")'
        sheet = load_workbook(BytesIO(workbook_bytes(block, state, date(2026, 10, 19)))).active
        dates = [row for row in sheet if row[0].is_date]
        self.assertEqual(len(dates), 24)
        self.assertEqual(dates[0][4].value, 60)
        self.assertTrue(all(row[4].value is None for row in dates[1:]))
        self.assertEqual(sheet['A5'].data_type, 's')
        self.assertEqual(sheet['A5'].value, block['weeks'][0])
        self.assertEqual(sheet.freeze_panes, 'C4')


class TelegramPlanPresentation(unittest.IsolatedAsyncioTestCase):
    async def test_default_plan_is_text_not_photo(self):
        block, state = fixture()
        message = SimpleNamespace(reply_text=AsyncMock(), reply_photo=AsyncMock())
        update = SimpleNamespace(message=message, effective_user=SimpleNamespace(id=1))
        with patch('nextset.coaching.programs.load', return_value=state):
            await ui.calendar(update, None)
        message.reply_photo.assert_not_called()
        self.assertEqual(message.reply_text.await_count, 5)
        self.assertIn('План Excel', str(message.reply_text.call_args))
